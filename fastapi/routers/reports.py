from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from database import get_db
import models
import os
import pandas as pd
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

router = APIRouter()

REPORTS_DIR = "/tmp/reports"
os.makedirs(REPORTS_DIR, exist_ok=True)

@router.get("/{formato}/{tipo}")
def generar_reporte(formato: str, tipo: str, db: Session = Depends(get_db)):
    if formato not in ['pdf', 'xlsx', 'docx']:
        raise HTTPException(status_code=400, detail="Invalid format. Use pdf, xlsx or docx")
        
    data = []
    titulo = "Reporte"
    
    if tipo == "ventas_globales":
        ordenes = db.query(models.Orden).all()
        data = [{"ID Orden": o.id, "Total": float(o.total), "Estado": o.estado, "Fecha": o.creado_en} for o in ordenes]
        titulo = "Reporte de Ventas Globales"
    elif tipo == "stock_bajo":
        inv = db.query(models.Inventario).filter(models.Inventario.stock_actual <= models.Inventario.stock_minimo).all()
        data = [{"ID Prod": i.producto_id, "Stock": i.stock_actual, "Mínimo": i.stock_minimo} for i in inv]
        titulo = "Reporte de Stock Bajo"
    elif tipo == "ordenes_pendientes":
        ordenes = db.query(models.Orden).filter(models.Orden.estado == 'Pendiente').all()
        data = [{"ID Orden": o.id, "Total": float(o.total), "Cliente ID": o.usuario_id} for o in ordenes]
        titulo = "Órdenes Pendientes de Envío"
    elif tipo == "logs_empleados":
        logs = db.query(models.LogMovimiento).all()
        data = [{"ID": l.id, "Usuario": l.usuario_id, "Acción": l.accion, "Fecha": l.fecha} for l in logs]
        titulo = "Registro de Actividades (Logs)"
    else:
         raise HTTPException(status_code=400, detail="Invalid report type")
         
    if not data:
        data = [{"Mensaje": "Sin datos"}]
         
    df = pd.DataFrame(data)
    filepath = f"{os.path.join(REPORTS_DIR, tipo)}.{formato}"
    
    if formato == 'xlsx':
        df.to_excel(filepath, index=False)
    elif formato == 'docx':
        doc = Document()
        doc.add_heading(titulo, 0)
        table = doc.add_table(rows=1, cols=len(df.columns))
        hdr_cells = table.rows[0].cells
        for i, col in enumerate(df.columns):
            hdr_cells[i].text = str(col)
        for _, row in df.iterrows():
            row_cells = table.add_row().cells
            for i, val in enumerate(row):
                row_cells[i].text = str(val)
        doc.save(filepath)
    elif formato == 'pdf':
        c = canvas.Canvas(filepath, pagesize=letter)
        c.drawString(100, 750, titulo)
        y = 700
        for _, row in df.iterrows():
            texto = " | ".join([f"{col}: {val}" for col, val in zip(df.columns, row)])
            c.drawString(50, y, texto)
            y -= 20
            if y < 50:
                c.showPage()
                y = 750
        c.save()
        
    return FileResponse(path=filepath, filename=f"Reporte_{tipo}.{formato}")
